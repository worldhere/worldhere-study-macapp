import datetime
import io
import openpyxl
from flask import Blueprint, request, jsonify, send_file

from db import get_db
from utils import parse_date, abs_min_to_datetime, abs_min_to_label, calc_working_minutes, format_elapsed
from models import load_shift_config, list_history_schedules
from auto_assign import auto_assign_tasks, mass_delay, _extend_end_over_breaks

bp = Blueprint('schedules', __name__)


@bp.route('/api/history_schedules')
def api_history_schedules():
    date_from = request.args.get("date_from") or ""
    date_to = request.args.get("date_to") or ""
    rows = list_history_schedules(
        date_from if date_from else None,
        date_to if date_to else None,
    )
    return jsonify({"history": rows})


def _bool(v, default=True):
    """Parse boolean from JSON, handling string 'false'/'0'/'' """
    if v is None:
        return default
    if isinstance(v, bool):
        return v
    if isinstance(v, str):
        return v.lower() not in ("false", "0", "")
    return bool(v)


def _parse_package_group_map(raw):
    """将前端传来的 {pkgId: groupName} 转为 {int(pkgId): groupName}"""
    if not raw or not isinstance(raw, dict):
        return None
    result = {}
    for k, v in raw.items():
        try:
            result[int(k)] = str(v)
        except (ValueError, TypeError):
            pass
    return result if result else None


def _parse_auto_assign_params(d):
    """Parse auto-assign request parameters, returns dict"""
    date = parse_date(d.get("date"))
    task_ids = d.get("task_ids") or []
    machine_ids = d.get("machine_ids") or []
    gap = int(d.get("gap", 0)) if d.get("gap") else 0
    ws = d.get("work_start_min")
    we = d.get("work_end_min")
    exclusion = d.get("exclusion_periods") or []
    exclusion = [(int(e[0]), int(e[1])) for e in exclusion if len(e) == 2]

    return {
        "date": date,
        "task_ids": task_ids,
        "machine_ids": machine_ids,
        "gap_minutes": max(0, min(120, gap)),
        "work_start_min": int(ws) if ws is not None else None,
        "work_end_min": int(we) if we is not None else None,
        "exclusion_periods": exclusion,
        "avoid_break_start": _bool(d.get("avoid_break_start"), False),
        "avoid_break_end": _bool(d.get("avoid_break_end"), False),
        "extend_over_breaks": _bool(d.get("extend_over_breaks"), True),
        "package_group_map": _parse_package_group_map(d.get("package_group_map")),
    }


@bp.route('/auto_assign_preview', methods=['POST'])
def auto_assign_preview():
    d = request.get_json() or {}
    p = _parse_auto_assign_params(d)
    result = auto_assign_tasks(
        task_ids=p["task_ids"],
        machine_ids=p["machine_ids"],
        date=p["date"],
        gap_minutes=p["gap_minutes"],
        work_start_min=p["work_start_min"],
        work_end_min=p["work_end_min"],
        exclusion_periods=p["exclusion_periods"],
        avoid_break_start=p["avoid_break_start"],
        avoid_break_end=p["avoid_break_end"],
        extend_over_breaks=p["extend_over_breaks"],
        package_group_map=p["package_group_map"],
        dry_run=True,
    )
    return jsonify(result)


@bp.route('/auto_assign', methods=['POST'])
def api_auto_assign():
    d = request.get_json() or {}
    p = _parse_auto_assign_params(d)
    result = auto_assign_tasks(
        task_ids=p["task_ids"],
        machine_ids=p["machine_ids"],
        date=p["date"],
        gap_minutes=p["gap_minutes"],
        work_start_min=p["work_start_min"],
        work_end_min=p["work_end_min"],
        exclusion_periods=p["exclusion_periods"],
        avoid_break_start=p["avoid_break_start"],
        avoid_break_end=p["avoid_break_end"],
        extend_over_breaks=p["extend_over_breaks"],
        package_group_map=p["package_group_map"],
        dry_run=False,
    )
    result["msg"] = f"分配完成：成功 {len(result['assigned'])} 个，未分配 {len(result['unassigned'])} 个"
    return jsonify(result)


@bp.route('/compact_tasks', methods=['POST'])
def api_compact_tasks():
    """回收后压实：将指定机器指定日期上 hole_start_min 之后的任务逐任务前移。
    如果提供了 now_min（当前时间，分钟），光标从 max(hole_start_min, now_min) 开始，
    防止任务被压实到已经过去的时间。"""
    d = request.get_json() or {}
    machine_id = int(d.get("machine_id", 0))
    date = d.get("date", "")
    hole_start_min = int(d.get("hole_start_min", 0))
    hole_end_min = int(d.get("hole_end_min", 0))
    now_min = d.get("now_min")
    now_min = int(now_min) if now_min is not None else None
    gap_minutes = max(0, min(120, int(d.get("gap_minutes", 0))))
    avoid_break_start = _bool(d.get("avoid_break_start"), False)
    avoid_break_end = _bool(d.get("avoid_break_end"), False)
    extend_over_breaks = _bool(d.get("extend_over_breaks"), True)

    if not machine_id or not date:
        return jsonify({"msg": "参数错误：需要 machine_id 和 date"}), 400

    from utils import compact_machine_tasks
    conn = get_db()
    shifted = compact_machine_tasks(
        conn, machine_id, date, hole_start_min, hole_end_min,
        gap_minutes=gap_minutes,
        avoid_break_start=avoid_break_start,
        avoid_break_end=avoid_break_end,
        extend_over_breaks=extend_over_breaks,
        now_min=now_min,
    )
    conn.commit()
    conn.close()
    return jsonify({"msg": f"已压实 {shifted} 个任务" if shifted else "无需压实", "shifted": shifted})


@bp.route('/mass_delay', methods=['POST'])
def api_mass_delay():
    d = request.get_json() or {}
    date = parse_date(d.get("date"))
    machine_ids = d.get("machine_ids") or []
    delay = int(d.get("delay_minutes", 0)) if d.get("delay_minutes") else 0
    from_min = int(d.get("from_start_min", 0)) if d.get("from_start_min") else 0
    mode = str(d.get("mode") or "shift").strip()
    strategy = str(d.get("strategy") or "block").strip()
    include_completed = bool(d.get("include_completed"))
    extend_over_breaks = d.get("extend_over_breaks", True)
    if mode not in ("shift", "extend"):
        mode = "shift"
    if strategy not in ("block", "smart"):
        strategy = "block"

    if not machine_ids:
        return jsonify({"msg": "请选择机器"}), 400
    if delay == 0:
        return jsonify({"msg": "延迟分钟数不能为0"}), 400

    result = mass_delay(
        machine_ids=[int(m) for m in machine_ids],
        date=date,
        delay_minutes=max(-28 * 1440, min(28 * 1440, delay)),
        from_start_min=max(0, min(28 * 1440, from_min)),
        mode=mode,
        strategy=strategy,
        include_completed=include_completed,
        extend_over_breaks=extend_over_breaks if isinstance(extend_over_breaks, bool) else True,
    )
    result["msg"] = f"已延迟 {result['affected']} 条排班记录"
    return jsonify(result)


@bp.route('/export_schedules', methods=['GET', 'POST'])
def export_schedules():
    if request.method == 'POST':
        d = request.get_json(silent=True) or {}
        status = d.get("status", "completed")
        date_from = d.get("date_from") or ""
        date_to = d.get("date_to") or ""
        columns = d.get("columns") or []
    else:
        status = request.args.get("status", "executing")
        date_from = request.args.get("date_from") or ""
        date_to = request.args.get("date_to") or ""
        columns = []

    if status not in ("executing", "completed"):
        return jsonify({"msg": "status 必须为 executing 或 completed"}), 400

    conn = get_db()
    shift_config = load_shift_config()
    sql = """
        SELECT s.id, s.date, s.machine_id, s.machine_name, s.task_name, s.task_type, s.task_kind,
               s.start_min, s.end_min, s.duration, s.status, s.remark,
               s.completed_at, s.actual_start_min, s.actual_end_min,
               s.estimated_window,
               t.rbp_task_id, t.priority, t.difficulty, t.scene,
               t.general_category, t.source_link, t.expected_count,
               t.collection_req_id, t.collection_req_type, t.est_mode,
               t.package_id
        FROM schedules s
        LEFT JOIN tasks t ON s.task_id = t.id
        WHERE s.status=?
    """
    params = [status]
    if date_from and date_to:
        sql += " AND julianday(s.date) + CAST(s.start_min AS REAL) / 1440.0 < julianday(?) + 1"
        params.append(date_to)
        sql += " AND julianday(s.date) + CAST(s.end_min AS REAL) / 1440.0 > julianday(?)"
        params.append(date_from)
    elif date_from:
        sql += " AND julianday(s.date) + CAST(s.end_min AS REAL) / 1440.0 > julianday(?)"
        params.append(date_from)
    elif date_to:
        sql += " AND julianday(s.date) + CAST(s.start_min AS REAL) / 1440.0 < julianday(?) + 1"
        params.append(date_to)
    sql += " ORDER BY s.date ASC, s.machine_id ASC, s.start_min ASC"

    rows = conn.execute(sql, params).fetchall()

    import datetime as _dt
    repair_data_map = {}
    for r in rows:
        sid = int(r["id"])
        mid = int(r["machine_id"] or 0)
        if mid:
            s_date = r["date"]
            s_start_min = int(r["start_min"])
            s_end_min = int(r["end_min"])
            s_date_dt = _dt.date.fromisoformat(s_date) if s_date else None
            if s_date_dt:
                s_start_dt = _dt.datetime.combine(s_date_dt, _dt.time(0, 0)) + _dt.timedelta(minutes=s_start_min)
                s_end_dt = _dt.datetime.combine(s_date_dt, _dt.time(0, 0)) + _dt.timedelta(minutes=s_end_min)
                rp_rows = conn.execute(
                    """SELECT id, start_datetime, end_datetime FROM repair_log
                       WHERE machine_id=? AND start_datetime < ? AND (end_datetime > ? OR end_datetime IS NULL)
                       ORDER BY start_datetime""",
                    (mid, s_end_dt.isoformat(), s_start_dt.isoformat()),
                ).fetchall()
                periods = []
                total_min = 0
                for rp in rp_rows:
                    rp_start = _dt.datetime.fromisoformat(rp["start_datetime"])
                    rp_end = _dt.datetime.fromisoformat(rp["end_datetime"]) if rp["end_datetime"] else None
                    o_start = max(s_start_dt, rp_start)
                    o_end = min(s_end_dt, rp_end) if rp_end else s_end_dt
                    dur = max(0, int((o_end - o_start).total_seconds() // 60))
                    total_min += dur
                    sl = f"{rp_start.month:02d}-{rp_start.day:02d} {rp_start.hour:02d}:{rp_start.minute:02d}"
                    el = f"{rp_end.month:02d}-{rp_end.day:02d} {rp_end.hour:02d}:{rp_end.minute:02d}" if rp_end else "至今"
                    periods.append(f"{sl}~{el}")
                repair_data_map[sid] = {
                    "duration": format_elapsed(total_min) if total_min > 0 else "",
                    "periods": "; ".join(periods) if periods else "",
                }

    # 构建 package_id -> package_name 映射
    _export_package_names = {}
    if rows:
        pkg_ids = set()
        for r in rows:
            pid = r["package_id"]
            if pid is not None:
                pkg_ids.add(int(pid))
        if pkg_ids:
            placeholders = ",".join("?" * len(pkg_ids))
            pkg_rows = conn.execute(
                f"SELECT id, name FROM task_packages WHERE id IN ({placeholders})",
                list(pkg_ids),
            ).fetchall()
            for pr in pkg_rows:
                _export_package_names[int(pr["id"])] = pr["name"] or ""

    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    label = "排班执行中" if status == "executing" else "排班的已完成任务"
    period = f"（{date_from}~{date_to}）" if date_from or date_to else ""
    ws.title = f"{label}{period}"[:31]

    status_map = {"executing": "执行中", "completed": "已完成"}

    EST_MODE_LABELS = {"blank": "不填", "direct": "直接预估", "calc": "计算预估"}

    ALL_COLUMNS = [
        # === 基本信息（7列）===
        ("date",            "排班日期",     lambda r: r["date"] or ""),
        ("completed_at",    "完成时间",     lambda r: r["completed_at"] or ""),
        ("task_name",       "任务名称",     lambda r: r["task_name"]),
        ("machine_name",    "机器名称",     lambda r: r["machine_name"]),
        ("task_type",       "机型",         lambda r: r["task_type"]),
        ("task_kind",       "任务类型",     lambda r: r["task_kind"]),
        ("status",          "状态",         lambda r: status_map.get(r["status"], r["status"])),
        # === 时间与时长（9列）===
        ("start_time",      "开始时间",     lambda r: abs_min_to_datetime(int(r["start_min"]), r["date"])),
        ("end_time",        "结束时间",     lambda r: abs_min_to_datetime(int(r["end_min"]), r["date"])),
        ("actual_start",    "实际开始",
         lambda r: abs_min_to_datetime(int(r["actual_start_min"]), r["date"]) if r["actual_start_min"] is not None else ""),
        ("actual_end",      "实际结束",
         lambda r: abs_min_to_datetime(int(r["actual_end_min"]), r["date"]) if r["actual_end_min"] is not None else ""),
        ("duration",        "预估时长",     lambda r: r["duration"] or ""),
        ("elapsed",         "排班时长",
         lambda r: format_elapsed(max(0, int(r["end_min"]) - int(r["start_min"])))),
        ("working",         "工作时长",
         lambda r: format_elapsed(calc_working_minutes(
             int(r["start_min"]), int(r["end_min"]), r["date"], shift_config))),
        ("est_mode",        "预估模式",
         lambda r: EST_MODE_LABELS.get(r["est_mode"], r["est_mode"] or "不填")),
        ("est_window",      "预估窗口",     lambda r: r["estimated_window"] or ""),
        # === 任务详情（11列）===
        ("priority",        "优先级",       lambda r: r["priority"] or ""),
        ("difficulty",      "难度",         lambda r: r["difficulty"] or ""),
        ("rbp_task_id",     "RBP数采任务ID", lambda r: r["rbp_task_id"] or ""),
        ("scene",           "场景",         lambda r: r["scene"] or ""),
        ("general_category","通用类别",     lambda r: r["general_category"] or ""),
        ("source_link",     "来源链接",     lambda r: r["source_link"] or ""),
        ("expected_count",  "预期采集量",
         lambda r: str(r["expected_count"]) if r["expected_count"] is not None else ""),
        ("collection_req_id","数采需求ID",  lambda r: r["collection_req_id"] or ""),
        ("collection_req_type","数采需求类型", lambda r: r["collection_req_type"] or ""),
        ("remark",          "备注",         lambda r: r["remark"] or ""),
        ("package_name",    "所属任务包",
         lambda r: _export_package_names.get(r["package_id"], "")),
        # === 维修相关（2列）===
        ("repair_duration", "维修时长",
         lambda r: repair_data_map.get(int(r["id"]), {}).get("duration", "")),
        ("repair_periods",  "维修时间段",
         lambda r: repair_data_map.get(int(r["id"]), {}).get("periods", "")),
    ]

    col_map = {c[0]: (c[1], c[2]) for c in ALL_COLUMNS}

    if columns:
        active = [(k, *col_map[k]) for k in columns if k in col_map]
    else:
        active = [(c[0], c[1], c[2]) for c in ALL_COLUMNS]

    headers = [label for (_, label, _) in active]
    ws.append(headers)

    for r in rows:
        row_data = [fn(r) for (_, _, fn) in active]
        ws.append(row_data)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    fname = f"{label}{period}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=fname,
    )
